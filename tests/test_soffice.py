"""DOCX -> PDF through LibreOffice on Linux (convert._soffice_to_pdf, platform_linux.soffice_*).

Most tests put a fake `soffice` on PATH. The `libreoffice` test runs the real one when it is
installed (on the dev machine: the snap) and is skipped otherwise."""
import json
import os
import sys
import time
from pathlib import Path

import pytest

from dragshift import convert

LINUX = sys.platform.startswith("linux")
pytestmark = pytest.mark.skipif(not LINUX, reason="LibreOffice path is Linux-only")

FAKE = r'''#!{python}
import json, os, subprocess, sys, time
args = sys.argv[1:]
mode = os.environ.get("FAKE_SOFFICE_MODE", "ok")
with open(os.environ["FAKE_SOFFICE_LOG"], "w") as f:
    json.dump(args, f)
if mode == "fail":
    print("Error: source file could not be loaded")
    sys.exit(1)
if mode == "hang":  # a forked child like soffice.bin, then never finish
    child = subprocess.Popen([sys.executable, "-c", "import time; time.sleep(600)"])
    with open(os.environ["FAKE_SOFFICE_LOG"] + ".child", "w") as f:
        f.write(str(child.pid))
    time.sleep(600)
if mode == "ok":
    import pymupdf
    outdir = args[args.index("--outdir") + 1]
    src = args[-1]
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "Hello from fake soffice")
    doc.save(os.path.join(outdir, os.path.splitext(os.path.basename(src))[0] + ".pdf"))
'''


def make_docx(path: Path) -> Path:
    from docx import Document
    from docx.shared import Inches
    from PIL import Image
    png = path.with_name("pic.png")
    Image.new("RGB", (120, 80), (255, 120, 80)).save(png)
    d = Document()
    d.add_heading("Hello heading", 1)
    d.add_paragraph("Hello world from DragShift.")
    t = d.add_table(rows=2, cols=2)
    for i, c in enumerate(t._cells):
        c.text = f"cell {i}"
    d.add_picture(str(png), width=Inches(1.5))
    d.save(path)
    png.unlink()
    return path


def pdf_text(path: Path) -> str:
    with convert.fitz.open(path) as doc:
        return "".join(p.get_text() for p in doc)


@pytest.fixture
def lx():
    from dragshift import platform_linux
    platform_linux.soffice_binary.cache_clear()
    yield platform_linux
    platform_linux.soffice_binary.cache_clear()


@pytest.fixture
def src(tmp_path) -> Path:
    d = tmp_path / "docs folder"
    d.mkdir()
    return make_docx(d / "Report ü.docx")


@pytest.fixture
def fake(tmp_path, monkeypatch, lx):
    """A fake `soffice` alone on PATH. Returns a function reading back its argv."""
    bindir = tmp_path / "bin"
    bindir.mkdir()
    exe = bindir / "soffice"
    exe.write_text(FAKE.replace("{python}", sys.executable))
    exe.chmod(0o755)
    log = tmp_path / "argv.json"
    monkeypatch.setenv("PATH", str(bindir))
    monkeypatch.setenv("FAKE_SOFFICE_LOG", str(log))
    monkeypatch.setattr(lx, "SNAP_LIBREOFFICE", str(tmp_path / "no-snap"))
    return lambda: json.loads(log.read_text())


def workroot_entries(lx):
    root = lx.soffice_workroot(lx.soffice_binary())
    return list(root.iterdir()) if root.exists() else []


# ---------------------------------------------------------------- conversion
def test_uses_libreoffice_when_installed(src, fake, lx, capsys):
    before = src.read_bytes()
    out = convert.convert(src, "docx", "PDF")
    assert out == [src.with_suffix(".pdf")]
    assert "Hello from fake soffice" in pdf_text(out[0])
    assert src.read_bytes() == before  # original untouched
    assert workroot_entries(lx) == []  # profile, input copy and output dir removed
    assert "LibreOffice" not in capsys.readouterr().out

    argv = fake()
    root = lx.soffice_workroot(lx.soffice_binary())
    assert argv[:5] == ["--headless", "--norestore", "--convert-to", "pdf", "--outdir"]
    assert Path(argv[5]).parent.parent == root
    profile = [a for a in argv if a.startswith("-env:UserInstallation=")]
    assert len(profile) == 1 and profile[0].startswith(f"-env:UserInstallation={root.as_uri()}/")
    # LibreOffice reads a private copy, never the original path
    assert argv[-1] != str(src) and Path(argv[-1]).name == src.name


def test_never_overwrites_an_existing_pdf(src, fake):
    src.with_suffix(".pdf").write_bytes(b"keep me")
    out = convert.convert(src, "docx", "PDF")[0]
    assert out.name == "Report ü (1).pdf"
    assert src.with_suffix(".pdf").read_bytes() == b"keep me"
    assert "Hello from fake soffice" in pdf_text(out)


@pytest.mark.parametrize("mode, reason", [("fail", "exit 1, no PDF: Error: source file could not be loaded"),
                                          ("empty", "exit 0, no PDF")])
def test_failure_falls_back_to_builtin_renderer(src, fake, lx, monkeypatch, capsys, mode, reason):
    monkeypatch.setenv("FAKE_SOFFICE_MODE", mode)
    out = convert.convert(src, "docx", "PDF")[0]
    text = pdf_text(out)
    assert "Hello world from DragShift." in text and "fake" not in text
    assert f"LibreOffice failed ({reason}), using the built-in renderer" in capsys.readouterr().out
    assert workroot_entries(lx) == []


def test_missing_binary_falls_back(src, lx, monkeypatch, tmp_path, capsys):
    monkeypatch.setenv("PATH", str(tmp_path / "empty"))
    monkeypatch.setattr(lx, "SNAP_LIBREOFFICE", str(tmp_path / "no-snap"))
    monkeypatch.setattr(lx.shutil, "which", lambda name: None)
    out = convert.convert(src, "docx", "PDF")[0]
    assert "Hello world from DragShift." in pdf_text(out)
    assert "LibreOffice not found" in capsys.readouterr().out


def test_timeout_kills_the_process_group(src, fake, lx, monkeypatch, tmp_path, capsys):
    monkeypatch.setenv("FAKE_SOFFICE_MODE", "hang")
    monkeypatch.setattr(convert, "SOFFICE_TIMEOUT", 3)
    t = time.monotonic()
    out = convert.convert(src, "docx", "PDF")[0]
    assert time.monotonic() - t < 15
    assert "Hello world from DragShift." in pdf_text(out)
    assert "timed out after 3 s" in capsys.readouterr().out
    child = int((tmp_path / "argv.json.child").read_text())
    for _ in range(50):  # killed; gone once its parent (init) reaps it
        if not Path(f"/proc/{child}").exists() or "Z" in Path(f"/proc/{child}/stat").read_text().split()[2]:
            break
        time.sleep(0.05)
    else:
        pytest.fail(f"soffice child {child} still running")
    assert workroot_entries(lx) == []


def test_not_used_off_linux(src, monkeypatch):
    calls = []
    monkeypatch.setattr(convert, "_soffice_to_pdf", lambda *a: calls.append(a) or True)
    monkeypatch.setattr(convert.sys, "platform", "win32")
    monkeypatch.setattr(convert, "_word_to_pdf", lambda s, o: False)
    convert.convert(src, "docx", "PDF")
    monkeypatch.setattr(convert.sys, "platform", "darwin")
    convert.convert(src, "docx", "PDF")
    assert calls == []


# ---------------------------------------------------------------- discovery
def test_probe_order_and_cache(lx, monkeypatch, tmp_path):
    found = {}
    monkeypatch.setattr(lx.shutil, "which", lambda name: found.get(name))
    snap = tmp_path / "libreoffice"
    monkeypatch.setattr(lx, "SNAP_LIBREOFFICE", str(snap))
    runs = []

    def run(cmd, **kw):
        runs.append(cmd)
        return type("R", (), {"returncode": 0 if flatpak_ok else 1})()
    monkeypatch.setattr(lx.subprocess, "run", run)

    def probe():
        lx.soffice_binary.cache_clear()
        return lx.soffice_binary()

    flatpak_ok = False
    assert probe() is None
    found["flatpak"] = "/usr/bin/flatpak"
    assert probe() is None and runs[-1] == ["/usr/bin/flatpak", "info", "org.libreoffice.LibreOffice"]
    flatpak_ok = True
    assert probe() == ("/usr/bin/flatpak", "run", "org.libreoffice.LibreOffice")
    snap.write_text("#!/bin/sh\n")
    snap.chmod(0o755)
    assert probe() == (str(snap),)
    found["libreoffice"] = "/usr/bin/libreoffice"
    assert probe() == ("/usr/bin/libreoffice",)
    found["soffice"] = "/usr/bin/soffice"
    assert probe() == ("/usr/bin/soffice",)

    n = len(runs)
    found.clear()
    assert lx.soffice_binary() == ("/usr/bin/soffice",)  # cached for the process
    assert len(runs) == n


def test_workroot(lx, monkeypatch, tmp_path):
    monkeypatch.setenv("HOME", str(tmp_path))
    # a snap can't use /tmp or ~/.cache: its own ~/snap/<name>/common
    assert lx.soffice_workroot(("/snap/bin/libreoffice",)) == \
        tmp_path / "snap" / "libreoffice" / "common" / "DragShift"
    monkeypatch.setenv("XDG_CACHE_HOME", str(tmp_path / "cache"))
    assert lx.soffice_workroot(("/usr/bin/soffice",)) == tmp_path / "cache" / "DragShift" / "soffice"
    monkeypatch.setenv("XDG_CACHE_HOME", "relative")
    assert lx.soffice_workroot(("/usr/bin/flatpak", "run", "x")) == \
        tmp_path / ".cache" / "DragShift" / "soffice"


# ---------------------------------------------------------------- the real thing
@pytest.mark.libreoffice
def test_real_libreoffice(src, lx, capsys):
    cmd = lx.soffice_binary()
    if not cmd:
        pytest.skip("LibreOffice is not installed")
    t = time.monotonic()
    out = convert.convert(src, "docx", "PDF")[0]
    took = time.monotonic() - t
    log = capsys.readouterr().out
    assert "LibreOffice" not in log, log  # no fallback
    with convert.fitz.open(out) as doc:
        assert doc.page_count >= 1
        assert "Hello" in doc[0].get_text()
        assert "cell 3" in doc[0].get_text()  # the table
        assert doc[0].get_images()  # the picture
        assert "LibreOffice" in (doc.metadata.get("producer") or "") + (doc.metadata.get("creator") or "")
    assert workroot_entries(lx) == []
    leftover = [p for p in Path("/proc").iterdir() if p.name.isdigit()
                and _cmdline(p).count(str(lx.soffice_workroot(cmd))) and "pytest" not in _cmdline(p)]
    assert leftover == [], [_cmdline(p) for p in leftover]
    print(f"\nreal LibreOffice {' '.join(cmd)}: {took:.1f} s", file=sys.__stderr__)


@pytest.mark.libreoffice
def test_real_libreoffice_timeout_leaves_nothing_running(src, lx, monkeypatch, capsys):
    cmd = lx.soffice_binary()
    if not cmd:
        pytest.skip("LibreOffice is not installed")
    monkeypatch.setattr(convert, "SOFFICE_TIMEOUT", 2)  # well before a snap finishes
    out = convert.convert(src, "docx", "PDF")[0]
    assert "timed out after 2 s" in capsys.readouterr().out
    assert "Hello world from DragShift." in pdf_text(out)  # built-in renderer
    root = str(lx.soffice_workroot(cmd))
    for _ in range(100):
        leftover = [p for p in Path("/proc").iterdir() if p.name.isdigit()
                    and root in _cmdline(p) and "pytest" not in _cmdline(p)]
        if not leftover:
            break
        time.sleep(0.1)
    assert leftover == [], [_cmdline(p) for p in leftover]
    assert workroot_entries(lx) == []


def _cmdline(p: Path) -> str:
    try:
        return (p / "cmdline").read_bytes().replace(b"\0", b" ").decode(errors="replace")
    except OSError:
        return ""
