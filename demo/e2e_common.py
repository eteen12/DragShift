"""What the end-to-end runners share, on any OS: the sample files, the output watcher and
the verifiers. Nothing here drives the desktop (no pyautogui, winkit, Quartz or GNOME), so
demo/e2e_windows.py, demo/e2e_mac.py and demo/e2e_linux.py all import it."""
from __future__ import annotations

import sys
import time
from pathlib import Path

from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
import dragshift.convert  # noqa: E402,F401  (registers HEIC/AVIF support for the verifiers)


# ------------------------------------------------------------------ sample files
def make_photo(path: Path, size=(2400, 1600), gps=True):
    import random
    random.seed(3)
    im = Image.new("RGB", size)
    d = ImageDraw.Draw(im)
    for y in range(size[1]):
        t = y / size[1]
        d.line([(0, y), (size[0], y)], fill=(int(255 - 120 * t), int(170 - 90 * t), int(90 + 110 * t)))
    for _ in range(250):
        x, y = random.randrange(size[0]), random.randrange(size[1])
        r = random.randrange(10, 90)
        d.ellipse((x, y, x + r, y + r), fill=tuple(random.randrange(256) for _ in range(3)))
    ex = Image.Exif()
    ex[0x010F] = "Apple"
    ex[0x0110] = "iPhone 15 Pro"
    if gps:
        ex[0x8825] = {1: "N", 2: (24.0, 33.0, 18.0), 3: "W", 4: (81.0, 47.0, 0.0)}
    im.save(path, quality=95, exif=ex)


def make_pdf(path: Path, img: Path, pages=3, title="Report"):
    import pymupdf
    doc = pymupdf.open()
    for i in range(pages):
        pg = doc.new_page()
        pg.insert_text((72, 80), f"{title} — page {i + 1}", fontsize=22)
        for j in range(12):
            pg.insert_text((72, 120 + j * 18), "Lorem ipsum dolor sit amet, consectetur adipiscing.", fontsize=11)
        pg.insert_image(pymupdf.Rect(72, 360, 420, 600), filename=str(img))
    doc.set_metadata({"author": "Jimmy", "title": title})
    doc.save(path)


def make_files(folder: Path, prefix=""):
    folder.mkdir(exist_ok=True)
    make_photo(folder / f"{prefix}Photo.jpg")
    make_photo(folder / f"{prefix}Second.jpg", size=(1200, 900), gps=False)
    Image.open(folder / f"{prefix}Photo.jpg").resize((800, 533)).save(folder / f"{prefix}Graphic.png")
    make_pdf(folder / f"{prefix}Report.pdf", folder / f"{prefix}Second.jpg")
    make_pdf(folder / f"{prefix}Appendix.pdf", folder / f"{prefix}Second.jpg", pages=2, title="Appendix")
    from docx import Document
    d = Document()
    d.add_heading("Letter", 1)
    d.add_paragraph("Hello from the DragShift end-to-end test.")
    d.save(folder / f"{prefix}Letter.docx")
    (folder / f"{prefix}Notes.txt").write_text("Shopping list\n- coffee\n- limes\n", encoding="utf-8")


def snapshot(folder: Path) -> set[str]:
    return {str(p.relative_to(folder)) for p in folder.rglob("*")}


def new_files(folder: Path, before: set[str], timeout=15) -> list[Path]:
    """New files (ignoring in-progress '.~*.part'), once their sizes stop changing."""
    end = time.time() + timeout

    def current():
        now = snapshot(folder) - before
        return {n: (folder / n).stat().st_size for n in now
                if (folder / n).is_file() and not Path(n).name.startswith(".~")}
    while time.time() < end:
        a = current()
        if a:
            time.sleep(0.8)
            b = current()
            if a == b:
                return sorted(folder / n for n in b)
        time.sleep(0.25)
    return []


SOURCES = ["Photo.jpg", "Second.jpg", "Graphic.png", "Report.pdf", "Appendix.pdf", "Letter.docx", "Notes.txt"]


# ------------------------------------------------------------------ verification helpers
def img_ok(p: Path, fmt=None):
    with Image.open(p) as im:
        im.load()
        if fmt:
            assert im.format == fmt, f"{p.name}: format {im.format} != {fmt}"
        return im.size


def pdf_pages(p: Path) -> int:
    import pymupdf
    with pymupdf.open(p) as d:
        return d.page_count


# ------------------------------------------------------------------ menu layout (from the app itself)
def ring_index(paths: list[str], ring: str, label: str) -> tuple[int, int, int]:
    from dragshift.app import DragShiftApp
    rings = DragShiftApp.build_rings(None, paths)
    names = list(rings)
    items = rings[ring]
    labels = [i.label for i in items]
    return names.index(ring), labels.index(label), len(items)


def safe(name: str) -> str:
    return "".join(c if c.isalnum() or c in "-_ ." else "_" for c in name).replace(" ", "_")


# ------------------------------------------------------------------ verifiers
def v_format(fmt, ext):
    def v(outs, srcs):
        o = outs[0]
        assert o.suffix.lower() == ext, f"extension {o.suffix}"
        if fmt == "PDF":
            assert pdf_pages(o) >= 1
            return f"{o.name} · {pdf_pages(o)} page(s)"
        if fmt in ("DOCX",):
            from docx import Document
            assert Document(o).paragraphs
            return o.name
        if fmt == "TXT":
            assert o.read_text(encoding="utf-8").strip()
            return o.name
        size = img_ok(o, None if fmt in ("HEIC", "AVIF") else fmt)
        return f"{o.name} · {size[0]}×{size[1]}"
    return v


def v_pages_to_images(n, ext):
    def v(outs, srcs):
        imgs = [o for o in outs if o.suffix.lower() == ext]
        assert len(imgs) == n, f"{len(imgs)} images, expected {n}"
        for o in imgs:
            img_ok(o)
        return f"{n} × {ext}"
    return v


def v_compress(limit):
    def v(outs, srcs):
        o = outs[0]
        s0, s1 = Path(srcs[0]).stat().st_size, o.stat().st_size
        assert s1 <= limit, f"{s1} bytes > {limit}"
        img_ok(o)
        return f"{s0 // 1024} KB → {s1 // 1024} KB ({(1 - s1 / s0) * 100:.0f}% smaller)"
    return v


def v_dims(fn):
    def v(outs, srcs):
        with Image.open(srcs[0]) as a:
            sa = a.size
        sb = img_ok(outs[0])
        assert fn(sa, sb), f"{sa} → {sb}"
        return f"{sa[0]}×{sa[1]} → {sb[0]}×{sb[1]}"
    return v


def v_exists(outs, srcs):
    img_ok(outs[0])
    return outs[0].name


def v_redact(outs, srcs):
    with Image.open(outs[0]) as im:
        im = im.convert("RGB")
        w, h = im.size
        px = im.getpixel((w // 2, h // 2))
    assert sum(px) < 40, f"centre pixel {px} not blacked out"
    return f"{outs[0].name} · centre blacked out"


def v_no_gps_keep_camera(outs, srcs):
    from dragshift import imaging
    s = imaging.metadata_summary(outs[0])
    assert "Location" not in s, "GPS still present"
    assert "Camera" in s, "camera info was removed"
    return f"GPS removed, kept {s['Camera']}"


def v_no_meta(outs, srcs):
    from dragshift import imaging
    n, gps = imaging.count_metadata(outs[0])
    assert n == 0 and not gps, f"{n} fields remain"
    return "0 metadata fields left"


def v_pdf_n(n):
    def v(outs, srcs):
        pdfs = [o for o in outs if o.suffix.lower() == ".pdf"]
        assert len(pdfs) == n, f"{len(pdfs)} PDFs"
        return f"{n} PDFs"
    return v


def v_merged(outs, srcs):
    total = sum(pdf_pages(Path(s)) for s in srcs)
    got = pdf_pages(outs[0])
    assert got == total, f"{got} pages, expected {total}"
    return f"{got} pages"


def v_rotated(outs, srcs):
    import pymupdf
    with pymupdf.open(outs[0]) as d:
        assert all(p.rotation == 90 for p in d)
    return "all pages at 90°"


def v_protected(outs, srcs):
    import pymupdf
    d = pymupdf.open(outs[0])
    assert d.needs_pass and d.authenticate("keywest2026")
    return "AES-256, password works"


def v_pdf_meta(outs, srcs):
    import pymupdf
    with pymupdf.open(outs[0]) as d:
        assert not (d.metadata or {}).get("author")
    return "author/title removed"


def v_pdf_compress(outs, srcs):
    s0, s1 = Path(srcs[0]).stat().st_size, outs[0].stat().st_size
    assert pdf_pages(outs[0]) == pdf_pages(Path(srcs[0]))
    return f"{s0 // 1024} KB → {s1 // 1024} KB"


def v_extract(outs, srcs):
    assert outs and all(o.parent.name.endswith("images") for o in outs)
    for o in outs:
        img_ok(o)
    return f"{len(outs)} image(s) in {outs[0].parent.name}/"


def v_combined(outs, srcs):
    assert pdf_pages(outs[0]) == len(srcs)
    return f"{len(srcs)} images → {pdf_pages(outs[0])}-page PDF"
